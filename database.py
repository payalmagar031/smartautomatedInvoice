from pathlib import Path
from openpyxl import Workbook, load_workbook
from datetime import datetime

DATA_DIR = Path(__file__).parent / "data"
DATA_DIR.mkdir(exist_ok=True)

FILES = {
    "products": DATA_DIR / "Products.xlsx",
    "customers": DATA_DIR / "Customers.xlsx",
    "sales": DATA_DIR / "Sales.xlsx",
    "details": DATA_DIR / "Invoice_Details.xlsx",
}

HEADERS = {
    "products": ["Product ID", "Product", "Category", "Price", "Stock"],
    "customers": ["Customer ID", "Customer Name", "Mobile", "Email", "Address"],
    "sales": ["Invoice No", "Date", "Customer", "Subtotal", "Discount", "GST", "Grand Total"],
    "details": ["Invoice No", "Date", "Product ID", "Product", "Category", "Qty", "Price", "Amount"],
}

def ensure_files():
    for key, path in FILES.items():
        if not path.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            ws.append(HEADERS[key])
            wb.save(path)

def rows(key):
    ensure_files()
    wb = load_workbook(FILES[key], data_only=True)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))

def append_row(key, row):
    ensure_files()
    wb = load_workbook(FILES[key])
    ws = wb.active
    ws.append(row)
    wb.save(FILES[key])

def replace_rows(key, data):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS[key])
    for row in data:
        ws.append(row)
    wb.save(FILES[key])

def next_id(key, prefix):
    nums = []
    for row in rows(key):
        val = str(row[0] or "")
        if val.startswith(prefix):
            try: nums.append(int(val[len(prefix):]))
            except ValueError: pass
    return f"{prefix}{(max(nums) if nums else 0)+1:04d}"

def next_invoice():
    return next_id("sales", "INV-")

def get_products():
    return rows("products")

def get_customers():
    return rows("customers")

def find_product(product_id):
    for r in get_products():
        if str(r[0]) == str(product_id):
            return r
    return None

def update_stock(product_id, qty):
    data = get_products()
    out = []
    for r in data:
        r = list(r)
        if str(r[0]) == str(product_id):
            r[4] = max(0, float(r[4] or 0) - float(qty))
            if r[4].is_integer(): r[4] = int(r[4])
        out.append(r)
    replace_rows("products", out)
